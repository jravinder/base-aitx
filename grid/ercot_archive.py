"""Build data/ercot_dam_hourly_2025_2026.csv from ERCOT report 13060.

Source: ERCOT MIS public report 13060, "Historical DAM Load Zone and Hub Prices"
(NP4-180-ER). No login. Listing:
  https://www.ercot.com/misapp/GetReports.do?reportTypeId=13060
Files used (fetched 2026-09-26):
  DAMLZHBSPP_2025.zip  doclookupId=1177667469  (2.0 MB, full year 2025)
  DAMLZHBSPP_2026.zip  doclookupId=1276779821  (1.5 MB, posted 2026-09-20, year to date)
Download: https://www.ercot.com/misdownload/servlets/mirDownload?doclookupId=<id>

Run:
  python3 grid/ercot_archive.py <DAMLZHBSPP_2025.xlsx> <DAMLZHBSPP_2026.xlsx>

Output columns: date (YYYY-MM-DD), he (hour ending 1..24), aen, north, houston ($/MWh).
DST fall-back repeated hour (flag Y) is dropped, so that day has 24 rows; the spring-forward
day has 23.
"""
import csv
import os
import sys

import openpyxl

ZONES = {"LZ_AEN": "aen", "LZ_NORTH": "north", "LZ_HOUSTON": "houston"}
OUT = os.path.join(os.path.dirname(__file__), "..", "data", "ercot_dam_hourly_2025_2026.csv")


def rows(path):
    wb = openpyxl.load_workbook(path, read_only=True)
    for name in wb.sheetnames:
        for r in wb[name].iter_rows(min_row=2, values_only=True):
            if not r or r[3] not in ZONES or r[2] == "Y":
                continue
            m, d, y = r[0].split("/")
            yield f"{y}-{m}-{d}", int(r[1].split(":")[0]), ZONES[r[3]], float(r[4])


def main(paths):
    table = {}
    for p in paths:
        for date, he, z, price in rows(p):
            table.setdefault((date, he), {})[z] = price
    with open(OUT, "w", newline="") as f:
        w = csv.writer(f)
        w.writerow(["date", "he", "aen", "north", "houston"])
        for (date, he) in sorted(table):
            v = table[(date, he)]
            w.writerow([date, he] + [f"{v[z]:g}" for z in ("aen", "north", "houston")])
    print(f"wrote {OUT}: {len(table)} hours, {table and min(table)[0]} to {max(table)[0]}")


if __name__ == "__main__":
    main(sys.argv[1:])
